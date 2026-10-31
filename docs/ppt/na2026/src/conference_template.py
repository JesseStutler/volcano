"""Use the supplied KubeCon NA 2026 PowerPoint masters and layout artwork.

The source template is 10 x 5.625 inches. Scale its native layout shapes to
13.333 x 7.5 inches so the established diagram dimensions stay unchanged.
"""
from pathlib import Path
import hashlib,json
from pptx import Presentation
from pptx.util import Inches

TEMPLATE_NAME = 'kubecon-na-2026-template.pptx'
COVER_LAYOUT = 0
SECTION_LAYOUT = 1
BODY_LAYOUT = 5


def create_presentation(assets):
    source = Path(assets) / TEMPLATE_NAME
    prs = Presentation(source)
    original_width, original_height = prs.slide_width, prs.slide_height
    width, height = Inches(13.333333), Inches(7.5)
    sx, sy = width / original_width, height / original_height
    # Remove sample slides but keep the official masters, layouts and image parts.
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    for master in prs.slide_masters:
        for owner in [master, *master.slide_layouts]:
            for shape in owner.shapes:
                # Scaling a group's outer transform scales its children together.
                left, top, sw, sh = shape.left, shape.top, shape.width, shape.height
                shape.left, shape.top = round(left*sx), round(top*sy)
                shape.width, shape.height = round(sw*sx), round(sh*sy)
            for tag in ['a:rPr', 'a:defRPr', 'a:endParaRPr']:
                for element in owner._element.xpath('.//' + tag + '[@sz]'):
                    element.set('sz', str(round(int(element.get('sz'))*sy)))
    prs.slide_width, prs.slide_height = width, height
    info = {
        'source': TEMPLATE_NAME,
        'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'source_url': 'https://github.com/JesseStutler/volcano/blob/docs/kubecon-na-2026-assets/docs/ppt/KubeCon%20CloudNativeCon%20North%20America%202026%20-%20Branded%20PowerPoint.pptx',
        'source_commit': '91168f97d6508109cab64e160e1981d913498ec5',
        'original_inches': [original_width/914400, original_height/914400],
        'output_inches': [width/914400, height/914400],
        'layouts': {'cover': 'TITLE', 'section': 'SECTION_HEADER', 'content': '1_Title and body'},
        'method': 'Native template masters, layout artwork and images retained; outer layout transforms scaled uniformly. Original template file unchanged.',
    }
    (Path(assets).parent/'template-provenance.json').write_text(json.dumps(info,indent=2)+'\n')
    return prs


def add_slide(prs, cover=False, section=False):
    index = COVER_LAYOUT if cover else SECTION_LAYOUT if section else BODY_LAYOUT
    slide = prs.slides.add_slide(prs.slide_layouts[index])
    # Use our existing native titles and figures, not the template's empty prompts.
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)
    return slide
