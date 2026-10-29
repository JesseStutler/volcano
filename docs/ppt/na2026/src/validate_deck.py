"""Validate generated artifacts without changing them. No cluster tests are run."""
from pathlib import Path
import json, math
import yaml
from lxml import etree
from PIL import ImageFont
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
ns={'s':'http://www.w3.org/2000/svg'}
manifest=json.loads((ROOT/'manifest.json').read_text())['slides']
routes=json.loads((ROOT/'geometry-audit.json').read_text())
errors=[]

def intersects(p,q,b):
    # Clip the line segment against the interior of a text ink bounding box.
    x0,y0,x1,y1=b;dx=q[0]-p[0];dy=q[1]-p[1];lo=0.;hi=1.
    for v,w in [(-dx,p[0]-x0),(dx,x1-p[0]),(-dy,p[1]-y0),(dy,y1-p[1])]:
        if abs(v)<1e-9:
            if w<0:return False
        else:
            t=w/v
            if v<0:lo=max(lo,t)
            else:hi=min(hi,t)
            if lo>hi:return False
    return True

for slide in manifest:
    if not slide['svg']:continue
    f=ROOT/slide['svg'];r=etree.parse(str(f)).getroot()
    width=float(r.get('width'));height=float(r.get('height'));boxes=[]
    for el in r.findall('s:text',ns):
        size=float(el.get('font-size'));family='Mono' if 'Mono' in el.get('font-family','') else 'Sans'
        weight='Bold' if el.get('font-weight')=='700' else 'Regular'
        font=ImageFont.truetype(f'/usr/share/fonts/truetype/liberation/Liberation{family}-{weight}.ttf',round(size*4))
        y=float(el.get('y'));parts=el.findall('s:tspan',ns)
        if not parts:parts=[el]
        for t in parts:
            text=t.text or '';x=float(t.get('x',el.get('x')))
            if t is not el:y+=float(t.get('dy',0))
            w=font.getlength(text)/4
            if el.get('text-anchor')=='middle':x-=w/2
            elif el.get('text-anchor')=='end':x-=w
            _,top,_,bottom=font.getbbox(text,anchor='ls');top=y+top/4;bottom=y+bottom/4
            if x< -1 or x+w>width+1 or top< -1 or bottom>height+1:errors.append((slide['page'],'text outside canvas',text))
            if text:boxes.append((x,top,x+w,bottom,text))
    for route in routes[str(slide['page'])]:
        pts=route['points']
        if route.get('kind')!='fabric':
            assert all(x==xx or y==yy for (x,y),(xx,yy) in zip(pts,pts[1:])),(slide['page'],pts)
        if route['arrow']:
            left,tip,right=route['head'];assert tip==pts[-1]
            assert abs(math.dist(left,tip)-math.dist(right,tip))<1e-6
            dx,dy=tip[0]-pts[-2][0],tip[1]-pts[-2][1]
            assert abs((right[0]-left[0])*dx+(right[1]-left[1])*dy)<1e-6
        for p,q in zip(pts,pts[1:]):
            for x0,y0,x1,y1,t in boxes:
                if intersects(p,q,(x0+.8,y0+.8,x1-.8,y1-.8)):
                    errors.append((slide['page'],'route crosses text',t,p,q))
    assert (ROOT/'png'/(f.stem+'.png')).exists()
for p in sorted((ROOT/'examples').glob('*.yaml')):
    yaml.safe_load(p.read_text())
    assert not any(isinstance(t,(yaml.tokens.FlowMappingStartToken,yaml.tokens.FlowSequenceStartToken)) for t in yaml.scan(p.read_text())),p
prs=Presentation(ROOT.parent/'Volcano_KubeCon_NA_2026.pptx')
assert len(prs.slides)==len(manifest)
for i,s in enumerate(prs.slides,1):
    assert s.notes_slide.notes_text_frame.text.strip(),i
    svg=manifest[i-1]['svg']
    if svg:
        expected=(ROOT/'png'/(Path(svg).stem+'.png')).read_bytes()
        assert any(getattr(sh,'shape_type',None)==13 and sh.image.blob==expected for sh in s.shapes),(i,'PPT does not embed the current PNG')
    for shape in s.shapes:
        assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+1 and shape.top+shape.height<=prs.slide_height+1,(i,shape.name)
if errors:
    print(json.dumps(errors,ensure_ascii=False,indent=2));raise SystemExit(1)
print(f'{len(manifest)} slides: text bounds, notes, image paths, route/text clearance, orthogonal arrows and symmetric arrowheads passed.')
print('All YAML excerpts parse and use block style. No benchmark result claims were tested.')
