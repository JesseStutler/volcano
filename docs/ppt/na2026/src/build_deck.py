"""Build Volcano's presentation with the official KubeCon NA 2026 template.
SVG is the editable diagram source; PNGs freeze diagram geometry in PowerPoint.
Dependencies: python-pptx, cairosvg, Pillow. Fonts: Liberation Sans / Mono.
"""
from pathlib import Path
from html import escape
from copy import deepcopy
import base64,json,re,sys
from lxml import etree
import cairosvg
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT=Path(__file__).resolve().parents[1]; PPT=ROOT.parent; REPO=PPT.parents[1]
ASSETS=ROOT/'assets'; SVG=ROOT/'svg'; PNG=ROOT/'png'
REUSE_SVG='--reuse-svg' in sys.argv
for d in [ASSETS,SVG,PNG]:d.mkdir(exist_ok=True)
C={'ink':'#17283b','muted':'#414f60','line':'#52677f','light':'#a6b5c6','pale':'#f2f5f9','blue':'#164f89','prefill':'#dce9f8','decode':'#dce5ee','green':'#146653','mint':'#deeee8','purple':'#62448d','lav':'#eae3f3','orange':'#c14421','amber':'#8b621c','sand':'#fff5e2','white':'#ffffff'}
FONT='/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'
BOLD='/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'
MONO='/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf'
def col(c):return C.get(c,c)
def wrap(text,width,size,bold=False):
 f=ImageFont.truetype(BOLD if bold else FONT,size*2)
 result=[]
 for para in text.split('\n'):
  line=''
  for word in para.split():
   new=(line+' '+word).strip()
   if line and f.getlength(new)>width*2:result.append(line);line=word
   else:line=new
  result.append(line)
 return result

class Art:
 def __init__(self,name,w=1480,h=560):self.name=name;self.w=w;self.h=h;self.bg=[];self.edges=[];self.front=[];self.routes=[];self.nodes={};self.textboxes=[]
 def rect(self,x,y,w,h,fill='white',stroke='light',rx=8,layer='front',dash=False):
  el=f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{rx}" fill="{col(fill)}" stroke="{col(stroke) if stroke else "none"}" stroke-width="2.2"'+(' stroke-dasharray="7 5"' if dash else '')+'/>'
  getattr(self,layer).append(el)
 def text(self,x,y,text,size=30,c='ink',bold=False,width=1400,anchor='start',mono=False):
  lines=wrap(text,width,size,bold);lh=size*1.28
  # Explicit line positions make SVG/PNG layout deterministic.
  self.front.append(f'<text x="{x}" y="{y}" font-family="{ "Liberation Mono" if mono else "Liberation Sans"}, Arial, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{col(c)}" text-anchor="{anchor}">'+''.join(f'<tspan x="{x}" dy="{0 if i==0 else lh}">{escape(t)}</tspan>' for i,t in enumerate(lines))+'</text>')
  self.textboxes.append((x,y,width,len(lines)*lh,text));return len(lines)*lh
 def panel(self,x,y,w,h,title=None,fill='pale',stroke='light'):
  self.rect(x,y,w,h,fill,stroke,layer='bg')
  if title:self.text(x+22,y+36,title,23,'muted',True,width=w-44)
 def box(self,id,x,y,w,h,text,fill='white',size=27,bold=False,c='ink',stroke='line'):
  self.rect(x,y,w,h,fill,stroke);self.nodes[id]=(x,y,w,h)
  lines=wrap(text,w-30,size,bold);start=y+h/2-(len(lines)-1)*size*.64+size*.33
  self.text(x+w/2,start,'\n'.join(lines),size,c,bold,w-30,'middle');return id
 def anchor(self,id,side):
  x,y,w,h=self.nodes[id];return {'top':(x+w/2,y),'bottom':(x+w/2,y+h),'left':(x,y+h/2),'right':(x+w,y+h/2)}[side]
 def path(self,points,arrow=True,c='line',dash=False):
  route={'points':points,'arrow':arrow,'color':col(c)}
  self.routes.append(route)
  d='M'+' L'.join(f'{x},{y}' for x,y in points)
  self.edges.append(f'<path d="{d}" fill="none" stroke="{col(c)}" stroke-width="2.5" stroke-linejoin="round"'+(' stroke-dasharray="7 5"' if dash else '')+'/>')
  if arrow:
   # Explicit symmetric chevron, with the same tip coordinate as the shaft.
   # No marker scaling/rotation is delegated to the SVG or PPT renderer.
   x,y=points[-1];px,py=points[-2];dx=x-px;dy=y-py
   length=(dx*dx+dy*dy)**.5
   assert length>=12, (self.name,points)
   ux,uy=dx/length,dy/length
   left=(x-10*ux-5*uy,y-10*uy+5*ux)
   right=(x-10*ux+5*uy,y-10*uy-5*ux)
   route['head']=[left,(x,y),right]
   self.edges.append(f'<path d="M{left[0]},{left[1]} L{x},{y} L{right[0]},{right[1]}" fill="none" stroke="{col(c)}" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round"/>')
 def connect(self,a,aside,b,bside,via=None,arrow=True,dash=False):
  p=self.anchor(a,aside);q=self.anchor(b,bside)
  # Leave a small gap at the target border; arrow tip points at the target.
  ox,oy={'top':(0,-4),'bottom':(0,4),'left':(-4,0),'right':(4,0)}[bside];end=(q[0]+ox,q[1]+oy)
  if via is None:
   if p[0]==q[0] or p[1]==q[1]:points=[p,end]
   elif aside in ['top','bottom']:
    mid=(p[1]+q[1])/2;points=[p,(p[0],mid),(q[0],mid),end]
   else:
    mid=(p[0]+q[0])/2;points=[p,(mid,p[1]),(mid,q[1]),end]
  else:points=[p]+via+[end]
  self.path(points,arrow,dash=dash)
 def badge(self,x,y,text,fill='mint',c='green',w=210):
  self.rect(x,y,w,35,fill,None,4);self.text(x+12,y+24,text,18,c,True,width=w-24)
 def image(self,x,y,w,h,path,opacity=1):
  p=Path(path);mime='image/svg+xml' if p.suffix=='.svg' else ('image/jpeg' if p.suffix in ['.jpg','.jpeg'] else 'image/png');uri='data:'+mime+';base64,'+base64.b64encode(p.read_bytes()).decode()
  self.front.append(f'<g opacity="{opacity}"><image x="{x}" y="{y}" width="{w}" height="{h}" preserveAspectRatio="xMidYMid meet" href="{uri}"/></g>')
 def save(self):
  markers=''.join(f'<marker id="arrow-{color.lstrip(chr(35))}" viewBox="0 0 12 12" refX="12" refY="6" markerWidth="12" markerHeight="12" markerUnits="userSpaceOnUse" orient="auto"><path d="M0 0 L12 6 L0 12 Z" fill="{color}"/></marker>' for color in sorted(set(C.values())))
  content=f'<svg xmlns="http://www.w3.org/2000/svg" width="{self.w}" height="{self.h}" viewBox="0 0 {self.w} {self.h}"><title>{escape(self.name)}</title><defs>{markers}</defs>'+''.join(self.bg+self.edges+self.front)+'</svg>'
  p=SVG/(self.name+'.svg')
  if REUSE_SVG and p.exists():content=p.read_text()
  else:p.write_text(content)
  cairosvg.svg2png(bytestring=content.encode(),write_to=str(PNG/(self.name+'.png')),output_width=self.w*2,output_height=self.h*2)
  return PNG/(self.name+'.png')

def replica(a,id,x,y,label,role='prefill',w=280):
 a.panel(x,y,w,132,fill='white',stroke='line');a.text(x+16,y+30,label,22,'ink',True,width=w-32)
 a.box(id+'-p0',x+16,y+58,(w-48)/2,50,'Pod 0',role,21);a.box(id+'-p1',x+32+(w-48)/2,y+58,(w-48)/2,50,'Pod 1',role,21)

def node(a,id,x,y,label,fill='white',w=110):a.box(id,x,y,w,48,label,fill,22)

# Brand assets: preserve original geometry; only opacity changes for watermark.
logo_source=REPO/'docs/images/volcano-logo-svg.svg'
if not logo_source.exists():logo_source=ASSETS/'volcano-logo.svg'
r=etree.parse(str(logo_source)).getroot();ns={'s':'http://www.w3.org/2000/svg'}
for e in r.xpath('.//s:foreignObject',namespaces=ns):e.getparent().remove(e)
r.set('width','1197');r.set('height','200');(ASSETS/'volcano-logo.svg').write_text('\n'.join(line.rstrip() for line in etree.tostring(r).decode().splitlines())+'\n')
cairosvg.svg2png(bytestring=etree.tostring(r),write_to=str(ASSETS/'volcano-logo.png'),output_width=1200)
# The official icon is the colored path from the horizontal logo.
icon=etree.Element('{http://www.w3.org/2000/svg}svg',nsmap={None:'http://www.w3.org/2000/svg'});icon.set('viewBox','0 0 82 40')
for e in r.xpath('.//s:linearGradient',namespaces=ns):icon.append(deepcopy(e))
for e in r.xpath('.//s:path[@class="st0"]',namespaces=ns):
 e=deepcopy(e);e.set('fill','url(#SVGID_1_)');icon.append(e)
(ASSETS/'volcano-icon.svg').write_text('\n'.join(line.rstrip() for line in etree.tostring(icon).decode().splitlines())+'\n')
# Technical diagrams: retain original EU / China figures where they carry the
# explanation, then add readable API excerpts as separate SVG text.
arts={};titles={};subtitles={};sources={};oldnotes={};chapters={};examples={}
REUSED=ASSETS/'reused'
def art(n,name,title,sub='',chapter=0,old=None,source=''):
 a=Art(f'{n:02d}-{name}');arts[n]=a;titles[n]=title;subtitles[n]=sub;chapters[n]=chapter;sources[n]=source
 if old:oldnotes[n]=old
 return a

def rule(a,x,y,w):a.path([(x,y),(x+w,y)],False,c='light')
def rawtext(a,x,y,t,size=26,c='ink',bold=False,mono=False):
 a.front.append(f'<text x="{x}" y="{y}" xml:space="preserve" font-family="{"Liberation Mono" if mono else "Liberation Sans"}" font-size="{size}" font-weight="{700 if bold else 400}" fill="{col(c)}">{escape(t)}</text>')

def code(a,x,y,w,h,label,body,size=25,highlights=(),filename=None):
 a.rect(x,y,w,h,'#f5f7fa','line',3)
 a.rect(x,y,w,47,'white',None,0)
 a.path([(x,y+47),(x+w,y+47)],False,'blue')
 a.text(x+20,y+32,label,24,'blue',True,width=w-40)
 lh=size*1.19
 for i,line in enumerate(body.strip('\n').splitlines()):
  yy=y+75+i*lh
  if ImageFont.truetype(MONO,size*2).getlength(line)/2>w-36:raise ValueError(f'Code line too wide: {a.name}: {line}')
  if yy>y+h-7:raise ValueError(f'Code overflows {a.name}: {label}')
  if any(k in line for k in highlights):
   a.rect(x+3,yy-size+2,w-6,lh,'#dce9f8',None,0)
   a.rect(x+3,yy-size+2,5,lh,'blue',None,0)
  rawtext(a,x+20,yy,line,size,'blue' if any(k in line for k in highlights) else 'ink',False,True)
 if filename:
  examples[filename]=body.strip()+'\n';(ROOT/'examples'/filename).write_text(body.strip()+'\n')

def original(a,file,x,y,w,h,crop=None):
 # SVG viewport crops the source image without resampling or rewriting it.
 p=REUSED/file
 if not crop:a.image(x,y,w,h,p);return
 from PIL import Image
 iw,ih=Image.open(p).size
 cx,cy,cw,ch=crop;uri='data:image/png;base64,'+base64.b64encode(p.read_bytes()).decode()
 a.front.append(f'<svg x="{x}" y="{y}" width="{w}" height="{h}" viewBox="{cx} {cy} {cw} {ch}" preserveAspectRatio="xMidYMid meet"><image width="{iw}" height="{ih}" href="{uri}"/></svg>')

def note(a,x,y,num,head,body,w=400):
 a.text(x,y,num,23,'blue',True);a.text(x+45,y,head,28,bold=True,width=w-45)
 a.text(x+45,y+43,body,25,'muted',width=w-45)

# Four divider slides are composed with native PowerPoint text below.
dividers={3:(1,'Unified scheduling','Different paths. Shared infrastructure.','How do schedulers coordinate node access?'),7:(2,'Multi-level gang\n& network topology','From available nodes to runnable workloads.','What must run together, and where should it fit?'),12:(3,'Ecosystem','Workload controllers carry scheduling intent.','How do roles and replicas reach the scheduler?'),17:(4,'Benchmark\nframework & progress','Measure behavior across workload shapes.','What can we reproduce, compare, and improve?')}

# 04: composition in approved_revision.py.
a=art(4,'workload-demands','Different Workloads, Shared Infrastructure','Training, inference, and agent applications put different demands on the same cluster.',1,2)

# 05: retain the original detailed architecture, not a replacement box summary.
a=art(5,'unified-architecture','A Coordinated Scheduling Framework','Specialized scheduling paths coordinate through NodeShard.',1,3,'China 2026 slide 7; same architecture family as EU 2026.')
original(a,'china-07-unified-architecture.png',0,0,970,550,(0,115,3129,2051))
note(a,1020,65,'01','Workload paths','Batch-oriented scheduling and a Pod-oriented fast path.',445)
note(a,1020,245,'02','Local coordination','Parallel workers use Binder conflict checks.',445)
note(a,1020,425,'03','Shared node access','NodeShard coordinates candidate-node pools.',445)

# 06: original controller plus an actual NodeShard excerpt.
a=art(6,'nodeshard','NodeShard: Desired Nodes and Current Use','Changing a shard coordinates future placement; it does not migrate running Pods.',1,4,'China 2026 slide 9; docs/design/agent-scheduler.md; shard API types.')
original(a,'china-09-dynamic-node-sharding.png',0,0,620,545)
code(a,690,0,790,500,'NodeShard · spec / status excerpt','''spec:
  nodesDesired:
  - node1
  - node2
status:
  nodesInUse:
  - node0
  - node1
  nodesToRemove:
  - node0
  nodesToAdd:
  - node2''',29,('nodesDesired','nodesInUse'),filename='nodeshard.yaml')
a.text(710,500,'Desired membership ≠ immediate use',29,'blue',True,width=740)
a.text(710,541,'Coordinators reconcile node access.',26,'muted',width=740)

# 08: the original full workload / topology figure is the main explanation.
a=art(8,'group-structure','PodGroup and SubGroup: Preserve Workload Structure','Gang requirements apply to the workload and to its complete subgroups.',2,6,'China 2026 slide 14. Figure reused unchanged.')
original(a,'china-14-podgroup-subgroup.png',0,0,1480,510)
a.text(0,552,'A SubGroup policy selects Pods; replica identity can split them into multiple SubGroups.',25,'blue',True,width=1470)

# 09: retain the physical fabric from the original, make the API precise and readable.
a=art(9,'hypernode-api','HyperNode: From Physical Fabric to Scheduling Domains','A parent HyperNode groups child HyperNodes; a leaf groups Kubernetes Nodes.',2,7,'China 2026 slide 16 (physical fabric); benchmark/scripts/create-hypernodes.sh.')
a.text(8,30,'PHYSICAL FABRIC',21,'muted',True)
original(a,'china-16-hypernode.png',0,55,445,470,(0,55,465,515))
code(a,500,0,460,510,'Parent · HyperNode.spec','''tier: 2
tierName: spine
members:
- type: HyperNode
  selector:
    exactMatch:
      name: rack-a''',30,('type: HyperNode','name: rack-a'),filename='hypernode-parent-spec.yaml')
code(a,995,0,485,510,'Leaf · HyperNode.spec','''tier: 1
tierName: leaf
members:
- type: Node
  selector:
    labelMatch:
      matchLabels:
        rack: rack-a''',30,('type: Node','rack: rack-a'),filename='hypernode-leaf-spec.yaml')
a.text(0,553,'The same hierarchy is available through the standalone HyperNode component.',26,'blue',True,width=1470)

# 10: reuse the actual network drawing, with the two topology levels in YAML.
a=art(10,'multilevel-placement','Place Groups at Different Topology Levels','An eight-Pod instance spans a larger domain; Prefill and Decode use different locality constraints.',2,8,'EU 2026 slide 16 / China 2026 slide 17; PodGroup API. Selector and Decode fragment below are explanatory excerpts.')
original(a,'china-17-multilevel-topology.png',0,0,900,510)
code(a,940,0,540,531,'PodGroup.spec · key fields','''minMember: 8
networkTopology:
  mode: hard
  highestTierName: core
subGroupPolicy:
- name: prefill
  subGroupSize: 4
  networkTopology:
    mode: hard
    highestTierName: leaf
- name: decode
  subGroupSize: 4
  networkTopology:
    mode: hard
    highestTierName: spine'''.replace('    mode: hard\n','').replace('  mode: hard\n',''),28,('highestTierName',),filename='multilevel-topology-spec.yaml')
a.text(0,554,'Named tiers: core (3) · spine (2) · leaf (1)',24,'blue',True,width=910)

# 11: composition in approved_revision.py.
a=art(11,'instance-isolation','Keep Each Instance Together, Place Instances Apart','ServingGroup → PodGroup: a Kthena scenario. Anti-affinity integration in Kthena is proposed.',2,9,'docs/user-guide/how_to_use_podgroup_anti_affinity.md; group-topology-affinity plugin, v1.16.')

# 13: composition in approved_revision.py.
a=art(13,'ecosystem','Next: RBG, Kubeflow, and LWS','Build on current integrations; carry finer group structure and placement requirements.',3,10,'Kthena docs; KubeRay #5249 / #5369; Grove #516; RBG gang-scheduling docs; Kubeflow Trainer Volcano plugin; LWS README / DisaggregatedSet. Roadmap intent supplied by maintainers.')

# 14: strong existing Kthena figure, kept in its original form.
a=art(14,'kthena-mapping','Kthena: From Serving Roles to Scheduling Groups','One ServingGroup maps to one PodGroup; each role replica forms a SubGroup.',3,6,'China 2026 slide 22; Kthena ModelServing gang scheduling guide.')
original(a,'china-22-kthena-gang.png',0,0,970,545)
code(a,1010,0,470,535,'ModelServing.spec · excerpt','''schedulerName: volcano
replicas: 1
template:
  gangPolicy:
    minRoleReplicas:
      prefill: 1
      decode: 3
  roles:
  - name: prefill
    replicas: 1
    workerReplicas: 1
  - name: decode
    replicas: 3
    workerReplicas: 1''',28,('minRoleReplicas','workerReplicas'),filename='kthena-modelserving-spec.yaml')
a.text(1025,540,'Pod templates omitted.',23,'muted',width=440)

# 15: composition in polish_figures.py.
a=art(15,'ray-replicas','KubeRay: A Replica Can Span Multiple Pods','The replica label separates complete groups inside one WorkerGroup.',3,11,'KubeRay #5249; #5369 open as of 2026-10-29. Volcano SubGroup API.')

# 16: composition in approved_revision.py.
a=art(16,'subgroup-next','Next: Relationships Between SubGroups','Keep cooperating groups local while separating replicas inside the same PodGroup.',2,12,'docs/design/group-topology-affinity.md (proposed SubGroup relationships); roadmap #6166.')

# 18: composition in polish_figures.py.
a=art(18,'benchmark-architecture','A Repeatable Scheduling Benchmark Framework','Generate scenarios, observe API events, and retain comparable measurements.',4,13,'benchmark/README.md; benchmark/scripts; benchmark/config; benchmark/testcases.')

# 19: composition in approved_revision.py.
a=art(19,'benchmark-progress','Compare Schedulers on the Same Workload','Repeat controlled runs to understand throughput, latency, and placement tradeoffs.',4,14,'benchmark/README.md; benchmark/Makefile; benchmark/scripts/collect-report.sh.')

# 20: composition in approved_revision.py.
a=art(20,'benchmark-measurement','Which Scenarios Fit Which Scheduler?','Expand tests across arrival patterns, group structure, and network topology.',4,15,'benchmark/scripts/collect-report.sh and audit-exporter event definitions. No performance results selected.')

# 21–22: typographic close in polish_figures.py.
a=art(21,'takeaways','What We Will Work On Next','Next steps in group placement, controller integration, and benchmarking.',4,16)
a=art(22,'questions','Questions & Discussion','Workload shapes, integration feedback, and reproducible observations are welcome.',0,17)
from polish_figures import polish
polish(arts, art, code, rawtext, C, ASSETS)
from approved_revision import revise
revise(arts, art, code, rawtext, original, C, ASSETS)
from light_edit import light_edit
removed_notes=light_edit(arts)
subtitles={n:re.sub(r";\s+([A-Za-z])",lambda m:". "+m.group(1).upper(),s) for n,s in subtitles.items()}
from speaker_labels import speaker_labels
speaker_labels(arts, subtitles, removed_notes, C)
from balanced_captions import balance_captions
balance_captions(arts, subtitles, removed_notes, titles)
from visual_finish import finish
finish(arts, C)
from grove_layout import layout_grove
layout_grove(arts)
from community_slides import add_community_slides
add_community_slides(arts, art, titles, ASSETS)
subtitles[22]=""
from approved_batch import apply_batch
order=apply_batch(arts, art, code, original, titles, subtitles, sources, ASSETS)
from mapping_examples import write_examples
write_examples(ROOT)
page_for={source_id:page for page,source_id in enumerate(order,1)}
for n,a in arts.items():
 a.name=f'{page_for[n]:02d}-'+a.name.split('-',1)[1]
 a.save()
# Remove only stale generated body images after a page-order change.
valid={a.name for a in arts.values()}
for folder in [SVG,PNG]:
 for path in folder.iterdir():
  if path.suffix in ['.svg','.png'] and path.stem not in valid:path.unlink()
(ROOT/'geometry-audit.json').write_text(json.dumps({str(page_for[n]):a.routes for n,a in arts.items()},indent=2)+'\n')

# Native PowerPoint title/profile/divider text. Source notes follow the remapped pages.
from conference_template import create_presentation, add_slide
prs=create_presentation(ASSETS)
def ptext(sl,x,y,w,h,text,size=26,bold=False,c='ink',align=None):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Liberation Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(col(c).lstrip('#'));p.space_after=Pt(5)
  if align:p.alignment=align
 return sh
chapter_names={0:'MAINTAINER TRACK',1:'UNIFIED SCHEDULING',2:'MULTI-LEVEL GANG & NETWORK TOPOLOGY',3:'ECOSYSTEM',4:'BENCHMARK FRAMEWORK & PROGRESS'}
agenda=['Unified scheduling','Multi-level gang & network topology','Ecosystem','Benchmark framework & progress']
from revision_notes import NOTES, write_notes_markdown
assert set(order) == set(NOTES)
manifest=[]
for page,n in enumerate(order,1):
 special=n==1 or n in dividers
 sl=add_slide(prs,cover=n==1,section=n in dividers)
 # The official title bar reserves its right side for the conference logo.
 if not special:
  heading='Speakers' if n==2 else titles[n]
  lines=wrap(heading,10.0*72,28,True)
  heading_h=.49*len(lines)
  ptext(sl,.455,(1.528-heading_h)/2,10.0,heading_h,'\n'.join(lines),28,True,'white')
 if n==1:
  titles[n]='Volcano for the Agentic AI Era: Unified Scheduling Across Training, Inference, and Agents'
  ptext(sl,.50,1.79,12.15,.66,'Volcano for the Agentic AI Era:',40,True,'white')
  ptext(sl,.52,2.57,11.95,1.08,'Unified Scheduling Across Training,\nInference, and Agents',29,False,'white')
  ptext(sl,.52,3.80,11.6,.41,'Ricardo Aravena  ·  Rainie Li',23,True,'white')
  ptext(sl,.52,4.20,11.8,.29,'Salt Lake City · November 9–12, 2026',15,False,'white')
 elif n==2:
  titles[n]='Speakers'
  for x,person,photo,role,detail,link in [(.80,'Ricardo Aravena','ricardo.jpeg','AI Infrastructure Lead · Nunchux AI','CNCF TOC Alternate (GB)\nCloud native and AI infrastructure','nunchux.ai · CNCF TOC'),(7.05,'Rainie Li','rainie.jpeg','Pinterest','Data processing infrastructure\nApache YuniKorn PMC member','Apache YuniKorn · rainieli')]:
   sl.shapes.add_picture(str(ASSETS/photo),Inches(x),Inches(2.00),width=Inches(1.75),height=Inches(1.75))
   ptext(sl,x,4.02,5.4,.5,person,30,True)
   ptext(sl,x,4.68,5.3,.43,role,22,False,'blue')
   ptext(sl,x,5.27,5.3,.85,detail,20,False,'muted')
   ptext(sl,x,6.35,5.3,.30,link,15,False,'blue')
 elif n in dividers:
  c,title,bridge,question=dividers[n];titles[n]=title.replace('\n',' ');chapters[n]=c
  ptext(sl,.65,1.58,2,.56,f'{c:02d}',31,True,'white')
  ptext(sl,.65,2.20,8.45,1.40,title,38,True,'white')
  ptext(sl,.68,3.78,8.35,.60,bridge,22,False,'white')
  for i,t in enumerate(agenda,1):ptext(sl,9.35,1.72+(i-1)*.64,3.30,.63,f'{i:02d}  {t}',15,i==c,'white' if i==c else '#bed1ed')
 else:
  if subtitles[n]:ptext(sl,.52,1.64,12.25,.39,subtitles[n],19,False,'blue')
  sl.shapes.add_picture(str(PNG/(arts[n].name+'.png')),Inches(.50),Inches(2.10 if subtitles[n] else 1.80),width=Inches(12.333333),height=Inches(4.666667))
 if not special:
  ptext(sl,.52,7.13,11.5,.22,'Volcano  |  KubeCon + CloudNativeCon North America 2026',9,False,'muted')
  ptext(sl,12.05,7.09,.75,.28,f'{page:02d} / {len(order)}',11,False,'muted',PP_ALIGN.RIGHT)
 sl.notes_slide.notes_text_frame.text=NOTES[n]
 manifest.append({'page':page,'source_id':n,'title':titles[n],'chapter':chapters.get(n,0),'svg':f'svg/{arts[n].name}.svg' if n in arts else None,'source':sources.get(n,'Native PowerPoint text'),'reference_set':oldnotes.get(n)})
prs.core_properties.title=titles[1];prs.core_properties.subject='KubeCon NA 2026 — revised technical presentation'
prs.core_properties.author='Ricardo Aravena; Rainie Li';prs.core_properties.comments='Official KubeCon NA 2026 template applied 2026-10-31. Retained China/EU figures and all English speaker notes.'
out=PPT/'Volcano_KubeCon_NA_2026.pptx';prs.save(out)
write_notes_markdown(ROOT, order, titles)
(ROOT/'manifest.json').write_text(json.dumps({'deck':out.name,'slides':manifest,'examples_are_excerpts':True,'benchmark':'Framework and ongoing validation; no invented results.','template':'template-provenance.json'},ensure_ascii=False,indent=2)+'\n')
print(out);print(f'{len(order)} slides; latest approved structure, ecosystem mappings, local captions and community close.')
